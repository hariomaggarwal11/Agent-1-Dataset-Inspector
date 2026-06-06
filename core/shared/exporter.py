"""Export utilities - PDF, DOCX, and JSON report generation."""

import json
import io
from datetime import datetime

import numpy as np
import pandas as pd


def export_json(report, filename=None):
    """Export inspection report as JSON.

    Parameters
    ----------
    report : dict
        Complete inspection report.
    filename : str, optional
        Output filename.

    Returns
    -------
    bytes
        JSON content as bytes.
    """
    # Create a serializable copy
    serializable = _make_serializable(report)

    content = json.dumps(serializable, indent=2, ensure_ascii=False)
    return content.encode("utf-8")


def _make_serializable(obj):
    """Convert report dict to JSON-serializable format."""
    if isinstance(obj, dict):
        return {k: _make_serializable(v) for k, v in obj.items()
                if k not in ("raw", "ecg_data")}
    elif isinstance(obj, pd.DataFrame):
        return obj.to_dict(orient="records")
    elif isinstance(obj, np.ndarray):
        return obj.tolist()
    elif isinstance(obj, (np.integer,)):
        return int(obj)
    elif isinstance(obj, (np.floating,)):
        return float(obj)
    elif isinstance(obj, (np.bool_,)):
        return bool(obj)
    elif isinstance(obj, list):
        return [_make_serializable(item) for item in obj]
    elif hasattr(obj, "__dict__"):
        return str(obj)
    else:
        return obj


def export_pdf(report, researcher_name=""):
    """Export inspection report as PDF.

    Parameters
    ----------
    report : dict
        Complete inspection report.
    researcher_name : str
        Researcher name for title page.

    Returns
    -------
    bytes
        PDF content as bytes.
    """
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.platypus import (
        SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak
    )
    from reportlab.lib import colors

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4,
                           leftMargin=2*cm, rightMargin=2*cm,
                           topMargin=2*cm, bottomMargin=2*cm)

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("CustomTitle", parent=styles["Title"],
                                 fontSize=18, spaceAfter=30)
    heading_style = ParagraphStyle("CustomHeading", parent=styles["Heading2"],
                                   fontSize=14, spaceAfter=12, spaceBefore=20)
    body_style = styles["Normal"]

    elements = []

    # Title page
    modality = report.get("modality", "Unknown")
    elements.append(Paragraph(f"{modality} Dataset Inspection Report", title_style))
    elements.append(Paragraph("NeuroInspect v1.0", styles["Heading3"]))
    elements.append(Spacer(1, 20))
    elements.append(Paragraph(f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}", body_style))
    if researcher_name:
        elements.append(Paragraph(f"Researcher: {researcher_name}", body_style))
    elements.append(Spacer(1, 40))

    # Metadata section
    elements.append(Paragraph("1. Dataset Overview", heading_style))
    metadata = report.get("metadata", {})
    if metadata:
        table_data = [["Field", "Value"]]
        for k, v in metadata.items():
            table_data.append([str(k), str(v)])
        t = Table(table_data, colWidths=[6*cm, 10*cm])
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.Color(0.04, 0.05, 0.1)),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 9),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.Color(0.95, 0.95, 0.95)]),
        ]))
        elements.append(t)

    # Quality Score
    quality = report.get("quality_score")
    if quality:
        elements.append(Spacer(1, 20))
        elements.append(Paragraph("2. Data Quality Score", heading_style))
        score = quality["score"]
        elements.append(Paragraph(f"Overall Score: {score}/100", styles["Heading3"]))
        elements.append(Spacer(1, 10))
        for item, pts in quality.get("breakdown", []):
            elements.append(Paragraph(f"  {item} ({pts})", body_style))

    # Artifacts
    artifacts = report.get("artifacts", {})
    if artifacts:
        elements.append(Spacer(1, 20))
        elements.append(Paragraph("3. Artifact Detection Results", heading_style))
        table_data = [["Check", "Status", "Details"]]
        for check_name, result in artifacts.items():
            if isinstance(result, dict):
                status = result.get("status", "unknown")
                detail = result.get("detail", "")
                table_data.append([check_name.replace("_", " ").title(), status.upper(), detail])
        if len(table_data) > 1:
            t = Table(table_data, colWidths=[5*cm, 3*cm, 8*cm])
            t.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.Color(0.04, 0.05, 0.1)),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, -1), 9),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
            ]))
            elements.append(t)

    doc.build(elements)
    return buffer.getvalue()


def export_docx(report, researcher_name=""):
    """Export inspection report as DOCX.

    Parameters
    ----------
    report : dict
        Complete inspection report.
    researcher_name : str
        Researcher name for title page.

    Returns
    -------
    bytes
        DOCX content as bytes.
    """
    from docx import Document
    from docx.shared import Pt, Cm
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document()

    # Set default font
    style = doc.styles["Normal"]
    font = style.font
    font.name = "Times New Roman"
    font.size = Pt(12)

    # Title
    modality = report.get("modality", "Unknown")
    title = doc.add_heading(f"{modality} Dataset Inspection Report", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph("NeuroInspect v1.0")
    doc.add_paragraph(f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
    if researcher_name:
        doc.add_paragraph(f"Researcher: {researcher_name}")
    doc.add_paragraph("")

    # Metadata
    doc.add_heading("1. Dataset Overview", level=1)
    metadata = report.get("metadata", {})
    if metadata:
        table = doc.add_table(rows=1, cols=2)
        table.style = "Table Grid"
        header = table.rows[0].cells
        header[0].text = "Field"
        header[1].text = "Value"
        for k, v in metadata.items():
            row = table.add_row().cells
            row[0].text = str(k)
            row[1].text = str(v)

    # Quality Score
    quality = report.get("quality_score")
    if quality:
        doc.add_heading("2. Data Quality Score", level=1)
        doc.add_paragraph(f"Overall Score: {quality['score']}/100")
        for item, pts in quality.get("breakdown", []):
            doc.add_paragraph(f"  {item} ({pts})", style="List Bullet")

    # Artifacts
    artifacts = report.get("artifacts", {})
    if artifacts:
        doc.add_heading("3. Artifact Detection Results", level=1)
        table = doc.add_table(rows=1, cols=3)
        table.style = "Table Grid"
        header = table.rows[0].cells
        header[0].text = "Check"
        header[1].text = "Status"
        header[2].text = "Details"
        for check_name, result in artifacts.items():
            if isinstance(result, dict):
                row = table.add_row().cells
                row[0].text = check_name.replace("_", " ").title()
                row[1].text = result.get("status", "unknown").upper()
                row[2].text = result.get("detail", "")

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
